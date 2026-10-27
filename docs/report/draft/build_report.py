"""Builds the project report (Word) from the chapter files in content/.

    python -I docs/report/draft/build_report.py

Needs python-docx and docxcompose. The front pages are the college's own
templates (template/1_ to 5_), filled in from the constants below; everything
after them uses the styles of the college report template
(template/7_report_styles.docx: CHNO, CHTITLE, HEAD1-3, MAINPARA, FIGLEGEND,
TABLE*). Writes output/Project_Report_Draft.docx. Word fills the contents and
the figure and table lists when the fields are updated (it asks on opening).

Chapter-file syntax, one block per blank-line-separated paragraph:
    # 1 | Introduction            chapter (CHNO + CHTITLE); "# - | Title" unnumbered,
                                  "# A | Title" appendix
    ## 1.1 Heading / ### / ####   HEAD1 / HEAD2 / HEAD3
    - item, 1. item               bullet and numbered lists
    ::table 6.1 | Caption         followed by a pipe table
    ::figure path | 4.1 | Caption | width-in-inches
    ::equation name | 4.1         docs/report_assets/equations/<name>.png
    ::algorithm 4.1 | Caption     followed by a ``` fenced block
    ::toc, ::lof, ::lot, ::references, ::pagebreak, ::frontend
Inline: **bold**, *italic*, `code`, [@K8] or [@K8, K26] (numbered in order of
first citation, IEEE style, from docs/report_prep/VERIFIED_REFERENCES.md).
"""

import copy
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docxcompose.composer import Composer

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
TEMPLATE = os.path.join(HERE, "template")
CONTENT = os.path.join(HERE, "content")
OUT = os.path.join(HERE, "output", "Project_Report_Draft.docx")
REFS = os.path.join(ROOT, "docs", "report_prep", "VERIFIED_REFERENCES.md")
EQUATIONS = os.path.join(ROOT, "docs", "report_assets", "equations")

TITLE = ("Intrusion Detection System using Hybrid CNN-LSTM with Explainable AI (SHAP) "
         "for Indian Network Databases")
STUDENTS = [("Maanya B S", "1JT23AI030"), ("Ruthwik Sai Ganesh", "1JT23AI042"), ("Varshini D N", "1JT23AI060")]
GUIDE = "Dr. Madhu B R"
PRINCIPAL = "Dr. K. Gopalakrishna"
ACADEMIC_YEAR = "2023-27"
RUNNING_TITLE = "Intrusion Detection System using Hybrid CNN-LSTM with Explainable AI"
KEYWORDS = ("intrusion detection, intrusion prevention, CNN-LSTM, deep learning, explainable AI, SHAP, "
            "CICIDS2017, MU-IoT, live packet capture.")
TEXT_WIDTH_IN = 6.0     # A4 minus the template's 1.25 in left and 1 in right margins

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|\[@[^\]]+\])")


# ---------------------------------------------------------------- references

def load_references():
    refs = {}
    with open(REFS, encoding="utf8") as f:
        for line in f:
            m = re.match(r"^(K\d+)\. (.+)$", line.strip())
            if m:
                text = re.sub(r"\s*\[(pages|page)[^\]]*\]\s*$", "", m.group(2))
                text = re.sub(r"\s*\(The CICFlowMeter reference\.\)\s*$", "", text)
                refs[m.group(1)] = text
    return refs


def citation_order(files):
    order = []
    for path in files:
        with open(path, encoding="utf8") as f:
            for group in re.findall(r"\[@([^\]]+)\]", f.read()):
                for key in (k.strip().lstrip("@") for k in group.split(",")):
                    if key not in order:
                        order.append(key)
    return {key: i + 1 for i, key in enumerate(order)}


# ---------------------------------------------------------------- low-level helpers

def text_runs(p):
    return [r for r in p.runs if not r._r.xpath(".//w:drawing|.//w:pict|.//w:object")]


def set_text(p, text):
    """Replaces a paragraph's text, keeping the formatting of its first text run."""
    runs = text_runs(p)
    if not runs:
        p.add_run(text)
        return
    keep = next((r for r in runs if r.text.strip()), runs[0])
    keep.text = text
    for r in runs:
        if r is not keep:
            r._r.getparent().remove(r._r)


def delete(p):
    p._p.getparent().remove(p._p)


def insert_after(p, text):
    new = copy.deepcopy(p._p)
    p._p.addnext(new)
    from docx.text.paragraph import Paragraph
    para = Paragraph(new, p._parent)
    set_text(para, text)
    return para


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    for tag, attrs in (("w:fldChar", {"w:fldCharType": "begin"}),):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        run._r.append(el)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run._r.append(instr)
    if placeholder:
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        run._r.append(sep)
        t = OxmlElement("w:t")
        t.text = placeholder
        run._r.append(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_inline(p, text, cites, size=None, bold=False, italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, mono = bold, italic, False
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, mono = part[1:-1], True
        elif part.startswith("[@"):
            keys = [k.strip().lstrip("@") for k in part[2:-1].split(",")]
            part = ", ".join(f"[{cites[k]}]" for k in keys)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, i = part[1:-1], True
        run = p.add_run(part)
        run.bold, run.italic = b or None, i or None
        if mono:
            run.font.name = "Courier New"
            run.font.size = Pt(10.5 if size is None else size - 1)
        elif size:
            run.font.size = Pt(size)
    return p


# ---------------------------------------------------------------- front pages

def fill_title_sheet(d):
    ps = d.paragraphs
    set_text(ps[6], f"“{TITLE}”")
    for p, (name, usn) in zip((ps[14], ps[15], ps[16]), STUDENTS):
        set_text(p, f"{usn}\t{name}")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent = Inches(1.35), Inches(0)
        pf.tab_stops.add_tab_stop(Inches(3.0))
    set_text(ps[20], GUIDE)
    set_text(ps[21], "Professor and Head")


def fill_certificate(d):
    names = ", ".join(f"{n} ({u})" for n, u in STUDENTS)
    set_text(d.paragraphs[3],
             f"This is to certify that the project work titled “{TITLE}” is carried out by {names}, "
             "bonafide students of Bachelor of Engineering at the Jyothy Institute of Technology, Bangalore, in "
             "partial fulfillment for the award of the degree of Bachelor of Engineering in Artificial "
             f"Intelligence and Machine Learning of Visvesvaraya Technological University, Belagavi, during the "
             f"academic year {ACADEMIC_YEAR}.")
    for row in d.tables[0].rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if "Dr./Prof." in p.text:
                    set_text(p, GUIDE)
                elif "Assistant/Associate Professor" in p.text:
                    set_text(p, p.text.replace("Assistant/Associate Professor", "Professor and Head (Guide)"))
    examiner = next(p for p in d.paragraphs if p.text.startswith("Name of the Examiner"))
    set_text(examiner, "Name of the Examiner\tSignature of the Examiner")
    examiner.paragraph_format.tab_stops.add_tab_stop(Inches(6.1), WD_TAB_ALIGNMENT.RIGHT)


def fill_declaration(d):
    set_text(d.paragraphs[0], "")
    names = ", ".join(f"{n} ({u})" for n, u in STUDENTS)
    set_text(d.paragraphs[2],
             f"We, {names}, are students of the eighth semester B.E. in Artificial Intelligence and Machine "
             "Learning at Jyothy Institute of Technology, VTU, and hereby declare that the project titled "
             f"“{TITLE}” has been carried out by us and submitted in partial fulfillment for the award "
             "of the degree of Bachelor of Engineering in Artificial Intelligence and Machine Learning during the "
             f"academic year {ACADEMIC_YEAR}. Further, the matter presented in the project has not been submitted "
             "previously by anybody for the award of any degree or any diploma to any other University, to the "
             "best of our knowledge and faith.")
    people = iter(STUDENTS)
    current = None
    for row in d.tables[0].rows:
        for cell in row.cells:
            for p in list(cell.paragraphs):
                t = p.text.strip()
                if t.startswith("Name"):
                    current = next(people, None)
                    set_text(p, f"Name: {current[0]}") if current else delete(p)
                elif t.startswith("USN"):
                    set_text(p, f"USN: {current[1]}") if current else delete(p)


def fill_acknowledgement(d):
    ps = d.paragraphs
    set_text(ps[4], f"In particular we would like to thank {PRINCIPAL}, Principal, Jyothy Institute of "
                    "Technology, VTU, for the constant encouragement and expert advice.")
    set_text(ps[7], f"We would like to thank our guide and Project Coordinator {GUIDE}, Professor and Head, "
                    "Department of Artificial Intelligence and Machine Learning, for sparing valuable time to extend "
                    "help in every step of our project work, and all the staff members of the Department of AIML "
                    "for their support.")
    set_text(ps[9], "We thank the University of New Brunswick and Manipur University for the CICIDS2017 and "
                    "MU-IoT datasets, and the owner of the network used for the live test for permission to use it.")
    # the template separates some paragraphs with empty ones; without them the page fits
    signature = next(p for p in d.paragraphs if "Signature of Students" in p.text)
    text_seen = False
    for p in list(d.paragraphs):
        if p is signature:
            break
        if p.text.strip():
            text_seen = True
        elif text_seen and not p._p.xpath(".//w:drawing|.//w:pict"):
            delete(p)
    for p in d.paragraphs:
        if p.text.strip() and p is not signature and p.text != "ACKNOWLEDGEMENT":
            p.paragraph_format.space_after = Pt(6)
    signature.paragraph_format.space_before = Pt(24)


def fill_abstract(d, paragraphs):
    first = d.paragraphs[1]
    set_text(first, paragraphs[0])
    last = first
    for text in paragraphs[1:]:
        last = insert_after(last, text)


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _style_value(style, get):
    while style is not None:
        v = get(style)
        if v is not None:
            return v
        style = style.base_style
    return None


def _doc_defaults(d):
    """Line spacing, spacing after/before (twips) and font size of the docDefaults."""
    el = d.styles.element.find(qn("w:docDefaults"))
    out = {"line": None, "after": None, "before": None, "size": None, "font": None}
    if el is None:
        return out
    sp = el.find(".//" + qn("w:pPrDefault") + "//" + qn("w:spacing"))
    if sp is not None:
        if sp.get(qn("w:line")):
            out["line"] = int(sp.get(qn("w:line"))) / 240
        for k in ("after", "before"):
            if sp.get(qn("w:" + k)):
                out[k] = Pt(int(sp.get(qn("w:" + k))) / 20)
    sz = el.find(".//" + qn("w:rPrDefault") + "//" + qn("w:sz"))
    if sz is not None:
        out["size"] = Pt(int(sz.get(qn("w:val"))) / 2)
    return out


def freeze_formatting(d):
    """Writes each paragraph's and run's formatting from its own document's styles as
    direct formatting, so that merging into the report (whose Normal style has 1.5 line
    spacing) does not change the college front pages' layout."""
    defaults = _doc_defaults(d)
    paragraphs = list(d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for p in paragraphs:
        style, pf = p.style, p.paragraph_format
        if pf.line_spacing is None:
            pf.line_spacing = _style_value(style, lambda s: s.paragraph_format.line_spacing) or defaults["line"] or 1.0
        if pf.space_after is None:
            pf.space_after = _style_value(style, lambda s: s.paragraph_format.space_after) or defaults["after"] or Pt(0)
        if pf.space_before is None:
            pf.space_before = _style_value(style, lambda s: s.paragraph_format.space_before) or defaults["before"] or Pt(0)
        if pf.alignment is None:
            align = _style_value(style, lambda s: s.paragraph_format.alignment)
            if align is not None:
                pf.alignment = align
        for r in p.runs:
            if r.font.size is None:
                r.font.size = _style_value(style, lambda s: s.font.size) or defaults["size"] or Pt(12)
            if r.font.name is None:
                r.font.name = _style_value(style, lambda s: s.font.name) or "Times New Roman"
            if r.font.bold is None:
                b = _style_value(style, lambda s: s.font.bold)
                if b:
                    r.font.bold = True


def format_abstract(d):
    """Template rule for the abstract: Times New Roman 12, justified, 1.5 line spacing."""
    for p in d.paragraphs[1:]:
        if not p.text.strip():
            continue
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing, pf.space_after, pf.space_before = 1.5, Pt(6), Pt(0)
        for r in p.runs:
            r.font.name, r.font.size = "Times New Roman", Pt(12)
        if p.text.startswith("Keywords:"):
            set_text(p, "")
            p.add_run("Keywords: ").bold = True
            p.add_run(KEYWORDS)
            for r in p.runs:
                r.font.name, r.font.size = "Times New Roman", Pt(12)


def add_toc_styles(styles):
    """Compact styles for the contents, figure and table lists (Word's 'toc 1'-'toc 3'):
    single spacing with a dotted right tab, instead of the body text's 1.5 lines."""
    from docx.oxml import parse_xml
    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    tab = int(TEXT_WIDTH_IN * 1440) - 10
    # level 4 (no indent, not bold) is used for the figure and table lists
    for level, indent, bold, before in ((1, 0, True, 80), (2, 280, False, 20), (3, 560, False, 0), (4, 0, False, 30)):
        sid = f"TOC{level}"
        old = styles.element.find(f'{qn("w:style")}[@{qn("w:styleId")}="{sid}"]')
        if old is not None:
            styles.element.remove(old)
        styles.element.append(parse_xml(
            f'<w:style {ns} w:type="paragraph" w:styleId="{sid}"><w:name w:val="toc {level}"/>'
            f'<w:basedOn w:val="Normal"/><w:next w:val="Normal"/><w:uiPriority w:val="39"/><w:unhideWhenUsed/>'
            f'<w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="{tab}"/></w:tabs>'
            f'<w:spacing w:before="{before}" w:after="20" w:line="264" w:lineRule="auto"/>'
            f'<w:ind w:left="{indent}"/><w:jc w:val="left"/></w:pPr>'
            f'<w:rPr>{"<w:b/>" if bold else ""}<w:spacing w:val="0"/></w:rPr></w:style>'))


def strip_trailing_empty(d):
    body = d.element.body
    for el in reversed(list(body)):
        if el.tag == qn("w:p") and not "".join(el.itertext()).strip() and not el.xpath(".//w:drawing|.//w:pict"):
            body.remove(el)
        elif el.tag == qn("w:sectPr"):
            continue
        else:
            break


def end_with_page_break(d):
    """Page break inside the last paragraph: a full page does not spill an empty one."""
    strip_trailing_empty(d)
    d.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- report body

class Builder:
    def __init__(self, doc, cites, refs):
        self.doc, self.cites, self.refs = doc, cites, refs
        self.counts = {"figures": 0, "tables": 0, "equations": 0, "words": 0}
        styles = doc.styles
        # The template's heading styles carry automatic list numbering; the chapter files
        # number headings themselves (and some headings are unnumbered), so it is removed.
        for sid in ("CHNO", "CHTITLE", "HEAD1", "HEAD2", "HEAD3"):
            el = styles.element.find(f'{qn("w:style")}[@{qn("w:styleId")}="{sid}"]')
            if el is not None:
                for num in el.findall(".//" + qn("w:numPr")):
                    num.getparent().remove(num)
        add_toc_styles(styles)
        if "TABLEGEND" not in [s.name for s in styles]:
            st = styles.add_style("TABLEGEND", 1)
            st.base_style = styles["FIGLEGEND"]
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.space_before = Pt(6)

    def para(self, text, style="MAINPARA", align=None):
        p = self.doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        add_inline(p, text, self.cites)
        self.counts["words"] += len(text.split())
        return p

    def chapter(self, num, title):
        if num == "-":
            p = self.doc.add_paragraph(style="CHTITLE")
            p.paragraph_format.page_break_before = True
            p.add_run(title)
            add_field(p, f'TC "{title}" \\l 1')
            return
        label = f"Appendix {num}" if num.isalpha() else f"Chapter {num}"
        p = self.doc.add_paragraph(style="CHNO")
        p.add_run(label)
        add_field(p, f'TC "{label}   {title}" \\l 1')
        self.doc.add_paragraph(title, style="CHTITLE")

    def heading(self, level, text):
        self.doc.add_paragraph(text, style={2: "HEAD1", 3: "HEAD2", 4: "HEAD3"}[level])

    def listing(self, items, numbered):
        for i, item in enumerate(items, 1):
            p = self.doc.add_paragraph(style="MAINPARA")
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Inches(0.35), Inches(-0.25)
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Inches(0.35))
            add_inline(p, (f"{i}.\t" if numbered else "•\t") + item, self.cites)
            self.counts["words"] += len(item.split())

    def caption(self, kind, number, text):
        style = "FIGLEGEND" if kind == "Fig." else "TABLEGEND"
        p = self.doc.add_paragraph(style=style)
        r = p.add_run(f"{kind} {number}. ")
        r.bold = True
        add_inline(p, text, self.cites)
        return p

    def table(self, number, caption, lines):
        rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
        aligns = None
        if len(rows) > 1 and all(re.fullmatch(r":?-{2,}:?", c) for c in rows[1]):
            aligns = ["right" if c.endswith(":") and not c.startswith(":") else
                      "left" if c.startswith(":") and not c.endswith(":") else None for c in rows[1]]
            rows = [rows[0]] + rows[2:]
        n = len(rows[0])
        if number != "-":                               # "::table - | ..." is an uncaptioned table
            self.caption("Table", number, caption)
        t = self.doc.add_table(rows=len(rows), cols=n)
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        size = 10 if number == "-" else 11 if n <= 4 else 10 if n <= 6 else 9
        weights = [max(4, min(38, max(len(re.sub(r"[*`]", "", r[j])) for r in rows))) for j in range(n)]
        widths = [TEXT_WIDTH_IN * w / sum(weights) for w in weights]
        for i, row in enumerate(rows):
            for j in range(n):
                cell = t.cell(i, j)
                cell.width = Inches(widths[j])
                p = cell.paragraphs[0]
                left = j == 0 or (aligns and aligns[j] == "left") or len(row[j]) > 28
                p.style = self.doc.styles["TABLELt" if left else "TABLECt"]
                p.paragraph_format.first_line_indent = Inches(0)
                # the template's table styles keep every row with the next, which pushes
                # whole tables onto a new page; only the header row is kept with the body
                p.paragraph_format.keep_with_next = i == 0
                add_inline(p, row[j] if j < len(row) else "", self.cites, size=size, bold=(i == 0))
                if i == 0:
                    shade(cell, "D9E2F3")
        trPr = t.rows[0]._tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader")
        hdr.set(qn("w:val"), "true")
        trPr.append(hdr)
        self.doc.add_paragraph(style="MAINPARA").paragraph_format.space_after = Pt(2)
        self.counts["tables"] += 1

    def figure(self, path, number, caption, width):
        full = os.path.join(ROOT, path)
        if not os.path.exists(full):
            raise FileNotFoundError(full)
        p = self.doc.add_paragraph(style="MAINPARA")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.first_line_indent = Inches(0)
        p.add_run().add_picture(full, width=Inches(float(width)))
        self.caption("Fig.", number, caption)
        self.counts["figures"] += 1

    def equation(self, name, number):
        full = os.path.join(EQUATIONS, name + ".png")
        if not os.path.exists(full):
            raise FileNotFoundError(full)
        from PIL import Image
        with Image.open(full) as im:
            w_px, h_px = im.size
        p = self.doc.add_paragraph(style="MAINPARA")
        pf = p.paragraph_format
        pf.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN / 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Inches(TEXT_WIDTH_IN), WD_TAB_ALIGNMENT.RIGHT)
        pf.space_before, pf.space_after = Pt(4), Pt(6)
        width_in = min(w_px / 300, TEXT_WIDTH_IN - 0.9)
        height_pt = h_px / w_px * width_in * 72
        p.add_run("\t")
        p.add_run().add_picture(full, width=Inches(width_in))
        num = p.add_run(f"\t({number})")
        # raise the number to the middle of the equation (w:position is in half-points)
        pos = OxmlElement("w:position")
        pos.set(qn("w:val"), str(max(0, int(height_pt - 9))))
        num._r.get_or_add_rPr().append(pos)
        self.counts["equations"] += 1

    def code(self, caption, lines):
        if caption:
            num, text = [s.strip() for s in caption.split("|", 1)]
            p = self.doc.add_paragraph(style="MAINPARA")
            p.paragraph_format.keep_with_next = True
            r = p.add_run(f"Algorithm {num}. ")
            r.bold = True
            p.add_run(text).bold = True
        t = self.doc.add_table(rows=1, cols=1)
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        shade(cell, "F2F2F2")
        first = True
        for ln in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.style = self.doc.styles["TABLELt"]
            pf = p.paragraph_format
            pf.space_after, pf.space_before, pf.line_spacing = Pt(0), Pt(0), 1.0
            r = p.add_run(ln.rstrip("\n") or " ")
            r.font.name, r.font.size = "Courier New", Pt(8.5)
        self.doc.add_paragraph(style="MAINPARA").paragraph_format.space_after = Pt(2)

    def special_heading(self, text):
        p = self.doc.add_paragraph(style="CHTITLE")
        p.paragraph_format.page_break_before = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(text)

    def field_page(self, title, instruction):
        self.special_heading(title)
        p = self.doc.add_paragraph(style="MAINPARA")
        add_field(p, instruction, "Update this field (right-click, Update Field) to fill this list.")

    def references(self):
        for key, n in sorted(self.cites.items(), key=lambda kv: kv[1]):
            if key not in self.refs:
                raise KeyError(f"citation {key} is not in {REFS}")
            p = self.doc.add_paragraph(style="MAINPARA")
            pf = p.paragraph_format
            pf.left_indent, pf.first_line_indent = Inches(0.45), Inches(-0.45)
            pf.tab_stops.add_tab_stop(Inches(0.45))
            pf.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, f"[{n}]\t" + self.refs[key], {})

    def render(self, path):
        with open(path, encoding="utf8") as f:
            lines = f.read().splitlines()
        i = 0
        while i < len(lines):
            ln = lines[i]
            s = ln.strip()
            if not s or s.startswith("<!--"):
                i += 1
                continue
            if s.startswith("::"):
                cmd, _, arg = s[2:].partition(" ")
                if cmd == "table":
                    num, cap = [x.strip() for x in arg.split("|", 1)]
                    j = i + 1
                    block = []
                    while j < len(lines) and lines[j].strip().startswith("|"):
                        block.append(lines[j])
                        j += 1
                    self.table(num, cap, block)
                    i = j
                    continue
                if cmd == "figure":
                    path_, num, cap, width = [x.strip() for x in arg.split("|")]
                    self.figure(path_, num, cap, width)
                elif cmd == "equation":
                    name, num = [x.strip() for x in arg.split("|")]
                    self.equation(name, num)
                elif cmd == "algorithm":
                    j = i + 2                          # next line must open a ``` fence
                    block = []
                    while not lines[j].strip().startswith("```"):
                        block.append(lines[j])
                        j += 1
                    self.code(arg, block)
                    i = j + 1
                    continue
                elif cmd == "toc":             # TOC \t takes style names (HEAD_1), not ids (HEAD1)
                    self.field_page("TABLE OF CONTENTS", 'TOC \\f \\h \\z \\t "HEAD_1,2,HEAD_2,3"')
                elif cmd == "lof":
                    self.field_page("LIST OF FIGURES", 'TOC \\h \\z \\t "FIG_LEGEND,4"')
                elif cmd == "lot":
                    self.field_page("LIST OF TABLES", 'TOC \\h \\z \\t "TABLEGEND,4"')
                elif cmd == "heading":
                    self.special_heading(arg)
                elif cmd == "references":
                    self.references()
                elif cmd == "pagebreak":
                    page_break(self.doc)
                elif cmd == "frontend":
                    end_front_section(self.doc)
                else:
                    raise ValueError(f"{path}: unknown directive {s}")
                i += 1
                continue
            if s.startswith("```"):
                j = i + 1
                block = []
                while not lines[j].strip().startswith("```"):
                    block.append(lines[j])
                    j += 1
                self.code(None, block)
                i = j + 1
                continue
            m = re.match(r"^(#{1,4}) (.+)$", s)
            if m:
                level, text = len(m.group(1)), m.group(2)
                if level == 1:
                    num, title = [x.strip() for x in text.split("|", 1)]
                    self.chapter(num, title)
                else:
                    self.heading(level, text)
                i += 1
                continue
            if re.match(r"^(- |\d+\. )", s):
                numbered = bool(re.match(r"^\d+\. ", s))
                items = []
                while i < len(lines) and lines[i].strip():
                    t = lines[i].strip()
                    if re.match(r"^(- |\d+\. )", t):
                        items.append(re.sub(r"^(- |\d+\. )", "", t))
                    else:
                        items[-1] += " " + t
                    i += 1
                self.listing(items, numbered)
                continue
            block = []
            while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith(("::", "#", "```", "|")):
                block.append(lines[i].strip())
                i += 1
            self.para(" ".join(block))


# ---------------------------------------------------------------- sections

def _page_numbers(sectPr, fmt, start=None):
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    cols = sectPr.find(qn("w:cols"))
    (cols.addprevious if cols is not None else sectPr.append)(pg)


def end_front_section(doc):
    """Ends the front-matter section: Roman numbers, none on the title sheet."""
    final = doc.element.body.find(qn("w:sectPr"))
    sect = copy.deepcopy(final)
    _page_numbers(sect, "lowerRoman", 1)
    if sect.find(qn("w:titlePg")) is None:
        tp = OxmlElement("w:titlePg")
        sect.find(qn("w:pgNumType")).addnext(tp)
    p = doc.add_paragraph()
    p._p.get_or_add_pPr().append(sect)


def finish_sections(doc):
    final = doc.element.body.find(qn("w:sectPr"))
    _page_numbers(final, "decimal", 1)
    for tp in final.findall(qn("w:titlePg")):
        final.remove(tp)
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)


# ---------------------------------------------------------------- main

def main():
    files = sorted(os.path.join(CONTENT, f) for f in os.listdir(CONTENT) if f.endswith(".md"))
    abstract_file = next(f for f in files if os.path.basename(f).startswith("00_abstract"))
    body_files = [f for f in files if f != abstract_file]
    cites = citation_order(body_files)
    refs = load_references()
    missing = [k for k in cites if k not in refs]
    if missing:
        sys.exit(f"Cited but not in VERIFIED_REFERENCES.md: {missing}")

    with open(abstract_file, encoding="utf8") as f:
        abstract = [p.strip().replace("\n", " ") for p in f.read().split("\n\n") if p.strip()
                    and not p.strip().startswith("<!--")]

    front = []
    for name, fill in (("1_title_sheet.docx", fill_title_sheet), ("2_certificate.docx", fill_certificate),
                       ("3_declaration.docx", fill_declaration), ("4_acknowledgement.docx", fill_acknowledgement)):
        d = Document(os.path.join(TEMPLATE, name))
        fill(d)
        freeze_formatting(d)
        end_with_page_break(d)
        front.append(d)
    d = Document(os.path.join(TEMPLATE, "5_abstract.docx"))
    fill_abstract(d, abstract)
    freeze_formatting(d)
    format_abstract(d)
    strip_trailing_empty(d)
    front.append(d)

    master = Document(os.path.join(TEMPLATE, "7_report_styles.docx"))
    body = master.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    composer = Composer(master)
    for d in front:
        composer.append(d)
    doc = composer.doc

    b = Builder(doc, cites, refs)
    for path in body_files:
        b.render(path)
    finish_sections(doc)
    for section in doc.sections:
        for p in section.header.paragraphs:
            if "Your Project Title" in p.text:
                set_text(p, p.text.replace("<Your Project Title>", RUNNING_TITLE))
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    unused = [k for k in refs if k not in cites]
    print(f"wrote {OUT}")
    print(f"body words ~{b.counts['words']:,}; figures {b.counts['figures']}; tables {b.counts['tables']}; "
          f"equations {b.counts['equations']}; references cited {len(cites)} (not cited: {unused or 'none'})")


if __name__ == "__main__":
    main()
